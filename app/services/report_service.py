from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


def create_excel_report(course, report):

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Attendance Report"

    # =====================================================
    # TITLE
    # =====================================================

    worksheet["A1"] = "ICT Attendance"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    # =====================================================
    # COURSE INFORMATION
    # =====================================================

    worksheet["A2"] = (
        f"{course.course_code} - "
        f"{course.course_name}"
    )

    worksheet["A2"].font = Font(
        bold=True
    )

    worksheet["A3"] = (
        f"Batch: "
        f"{course.batch.name if course.batch else 'N/A'}"
    )

    worksheet["A4"] = (
        f"Academic Session: "
        f"{course.academic_session or 'N/A'}"
    )

    # =====================================================
    # HEADERS
    # =====================================================

    headers = [
        "Student ID",
        "Student Name",
        "Batch",
        "Obtained Weight",
        "Total Weight",
        "Attendance %"
    ]

    header_row = 6

    for column_number, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column_number
        )

        cell.value = header
        cell.font = Font(
            bold=True
        )

    # =====================================================
    # DATA
    # =====================================================

    row_number = header_row + 1

    for item in report:

        worksheet.cell(
            row=row_number,
            column=1
        ).value = item["student_id"]

        worksheet.cell(
            row=row_number,
            column=2
        ).value = item["name"]

        worksheet.cell(
            row=row_number,
            column=3
        ).value = item["batch"]

        worksheet.cell(
            row=row_number,
            column=4
        ).value = item["obtained"]

        worksheet.cell(
            row=row_number,
            column=5
        ).value = item["total"]

        percentage_cell = worksheet.cell(
            row=row_number,
            column=6
        )

        percentage_cell.value = (
            item["percentage"] / 100
        )

        percentage_cell.number_format = "0.00%"

        row_number += 1

    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 30
    worksheet.column_dimensions["C"].width = 15
    worksheet.column_dimensions["D"].width = 20
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 18

    # =====================================================
    # SAVE TO MEMORY
    # =====================================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output
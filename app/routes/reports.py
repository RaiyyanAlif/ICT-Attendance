from io import BytesIO

from flask import Blueprint, render_template, send_file
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from app import db
from app.models.course import Course
from app.models.enrollment import Enrollment
from app.models.attendance import AttendanceSession, AttendanceRecord


reports_bp = Blueprint("reports", __name__, url_prefix="/reports")


def get_teacher_course(course_id):
    course = db.session.get(Course, course_id)

    if not course:
        return None

    if course.teacher_id != current_user.id:
        return None

    return course


@reports_bp.route("/course/<int:course_id>")
@login_required
def course_report(course_id):
    course = get_teacher_course(course_id)

    if not course:
        return "Course not found or unauthorized", 404

    enrollments = (
        Enrollment.query
        .filter_by(course_id=course.id)
        .all()
    )

    students = sorted(
        [enrollment.student for enrollment in enrollments],
        key=lambda student: student.student_id
    )

    sessions = (
        AttendanceSession.query
        .filter_by(course_id=course.id)
        .order_by(
            AttendanceSession.date.asc(),
            AttendanceSession.id.asc()
        )
        .all()
    )

    records = AttendanceRecord.query.filter(
        AttendanceRecord.session_id.in_(
            [session.id for session in sessions]
        )
    ).all() if sessions else []

    record_map = {
        (record.session_id, record.student_id): record.present
        for record in records
    }

    report_rows = []

    for student in students:
        total_weight = 0
        present_weight = 0

        for session in sessions:
            weight = session.attendance_weight or 1

            total_weight += weight

            if record_map.get((session.id, student.id), False):
                present_weight += weight

        percentage = (
            (present_weight / total_weight) * 100
            if total_weight > 0
            else 0
        )

        report_rows.append({
            "student": student,
            "present_weight": present_weight,
            "total_weight": total_weight,
            "percentage": round(percentage, 2)
        })

    # Batch session is the source of truth.
    academic_session = (
        course.batch.academic_session
        if course.batch
        else course.academic_session
    )

    return render_template(
        "reports/course_report.html",
        course=course,
        students=students,
        sessions=sessions,
        report_rows=report_rows,
        academic_session=academic_session
    )


@reports_bp.route("/course/<int:course_id>/excel")
@login_required
def course_excel(course_id):
    course = get_teacher_course(course_id)

    if not course:
        return "Course not found or unauthorized", 404

    enrollments = (
        Enrollment.query
        .filter_by(course_id=course.id)
        .all()
    )

    students = sorted(
        [enrollment.student for enrollment in enrollments],
        key=lambda student: student.student_id
    )

    sessions = (
        AttendanceSession.query
        .filter_by(course_id=course.id)
        .order_by(
            AttendanceSession.date.asc(),
            AttendanceSession.id.asc()
        )
        .all()
    )

    records = AttendanceRecord.query.filter(
        AttendanceRecord.session_id.in_(
            [session.id for session in sessions]
        )
    ).all() if sessions else []

    record_map = {
        (record.session_id, record.student_id): record.present
        for record in records
    }

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Attendance Report"

    academic_session = (
        course.batch.academic_session
        if course.batch
        else course.academic_session
    )

    # ---------------------------------------------------------
    # Course information
    # ---------------------------------------------------------

    worksheet["A1"] = "ICT Attendance - Attendance Report"
    worksheet["A1"].font = Font(bold=True, size=16)

    worksheet["A2"] = "Course Code"
    worksheet["B2"] = course.course_code

    worksheet["A3"] = "Course Name"
    worksheet["B3"] = course.course_name

    worksheet["A4"] = "Batch"
    worksheet["B4"] = (
        course.batch.name
        if course.batch
        else "Not Assigned"
    )

    worksheet["A5"] = "Academic Session"
    worksheet["B5"] = academic_session or "Not Set"

    worksheet["A6"] = "Teacher"
    worksheet["B6"] = current_user.name

    # ---------------------------------------------------------
    # Attendance table
    # ---------------------------------------------------------

    header_row = 8

    worksheet.cell(header_row, 1, "Student ID")
    worksheet.cell(header_row, 2, "Student Name")

    column = 3

    for session in sessions:
        header = (
            f"{session.date.strftime('%Y-%m-%d')} "
            f"(W{session.attendance_weight})"
        )

        worksheet.cell(header_row, column, header)
        column += 1

    worksheet.cell(header_row, column, "Present Weight")
    worksheet.cell(header_row, column + 1, "Total Weight")
    worksheet.cell(header_row, column + 2, "Attendance %")

    for cell in worksheet[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ---------------------------------------------------------
    # Student rows
    # ---------------------------------------------------------

    row = header_row + 1

    for student in students:

        worksheet.cell(row, 1, student.student_id)
        worksheet.cell(row, 2, student.name)

        current_column = 3

        present_weight = 0
        total_weight = 0

        for session in sessions:
            weight = session.attendance_weight or 1

            total_weight += weight

            present = record_map.get(
                (session.id, student.id),
                False
            )

            if present:
                worksheet.cell(
                    row,
                    current_column,
                    "Present"
                )

                present_weight += weight

            else:
                worksheet.cell(
                    row,
                    current_column,
                    "Absent"
                )

            current_column += 1

        percentage = (
            (present_weight / total_weight) * 100
            if total_weight > 0
            else 0
        )

        worksheet.cell(
            row,
            current_column,
            present_weight
        )

        worksheet.cell(
            row,
            current_column + 1,
            total_weight
        )

        worksheet.cell(
            row,
            current_column + 2,
            round(percentage, 2)
        )

        row += 1

    # ---------------------------------------------------------
    # Formatting
    # ---------------------------------------------------------

    for column_cells in worksheet.columns:
        max_length = 0

        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            try:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
            except Exception:
                pass

        worksheet.column_dimensions[column_letter].width = (
            min(max(max_length + 2, 12), 30)
        )

    worksheet.freeze_panes = "C9"

    # ---------------------------------------------------------
    # Send Excel file
    # ---------------------------------------------------------

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    filename = (
        f"{course.course_code}_"
        f"{course.batch.name if course.batch else 'course'}_"
        f"attendance.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# -------------------------------------------------------------
# Compatibility route
#
# Your current details.html is calling:
# reports.download_excel
#
# Keep this route so the existing template works.
# -------------------------------------------------------------

@reports_bp.route("/course/<int:course_id>/download-excel")
@login_required
def download_excel(course_id):
    return course_excel(course_id)